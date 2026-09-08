#!/usr/bin/env python3
"""Markdown → A4 Word (.docx) with the Poly navy spec layout."""
from __future__ import annotations

import argparse
import json
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

import markdown
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

NAVY = "1A365D"
TEAL = "2B6CB0"
ROW_ALT = "EDF2F7"
RULE = "CBD5E0"
BODY = RGBColor(0x1A, 0x20, 0x2C)
MUTED = RGBColor(0x4A, 0x55, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_RGB = RGBColor(0x1A, 0x36, 0x5D)
TEAL_RGB = RGBColor(0x2B, 0x6C, 0xB0)
CODE_FG = RGBColor(0x2D, 0x37, 0x48)
GRID = "A0AEC0"

FONT = "Calibri"
FONT_MONO = "Consolas"
USABLE_MM = 174.0


def set_run_font(run, name=FONT, size=10.5, bold=False, italic=False, color=BODY, mono=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = FONT_MONO if mono else name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    face = FONT_MONO if mono else name
    rFonts.set(qn("w:ascii"), face)
    rFonts.set(qn("w:hAnsi"), face)
    rFonts.set(qn("w:cs"), face)
    rFonts.set(qn("w:eastAsia"), face)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    tcPr.append(shd)


def set_cell_borders(cell, color=GRID, sz="4") -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcBorders"):
            tcPr.remove(child)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_cell_margins(cell, top=40, bottom=40, left=60, right=60) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcMar"):
            tcPr.remove(child)
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def set_cell_valign(cell, val="top") -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:vAlign"):
            tcPr.remove(child)
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), val)
    tcPr.append(vAlign)


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_table_fixed(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def set_table_width(table, width_mm: float) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    for child in list(tblPr):
        if child.tag == qn("w:tblW"):
            tblPr.remove(child)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(width_mm * 56.7)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)


def disable_autofit(table) -> None:
    table.autofit = False
    table.allow_autofit = False


def set_paragraph_spacing(p, before=0, after=6, line=240, left=0) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line / 240.0
    if left:
        pf.left_indent = Mm(left)


def add_page_field(paragraph, instruction: str = "PAGE") -> None:
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = f" {instruction} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin)
    r.append(instr)
    r.append(end)
    set_run_font(run, size=8, color=WHITE)


def add_inline_runs(paragraph, html: str, size=10.5, color=BODY, justify=False) -> None:
    tokens = re.split(r"(</?b>|</?i>|<code>|</code>|<br\s*/?>)", html)
    bold = italic = code = False
    for tok in tokens:
        if not tok:
            continue
        if tok == "<b>":
            bold = True
            continue
        if tok == "</b>":
            bold = False
            continue
        if tok == "<i>":
            italic = True
            continue
        if tok == "</i>":
            italic = False
            continue
        if tok == "<code>":
            code = True
            continue
        if tok == "</code>":
            code = False
            continue
        if tok.startswith("<br"):
            paragraph.add_run().add_break()
            continue
        text = unescape(tok)
        if not text:
            continue
        run = paragraph.add_run(text)
        set_run_font(
            run,
            size=8.5 if code else size,
            bold=bold,
            italic=italic,
            color=CODE_FG if code else color,
            mono=code,
        )


def unescape(text: str) -> str:
    return (
        text.replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&amp;", "&")
        .replace("&nbsp;", " ")
        .replace("&#39;", "'")
        .replace("&quot;", '"')
    )


def looks_like_meta(text: str) -> bool:
    plain = re.sub(r"</?b>", "", text)
    return (
        plain.startswith("Criado em:")
        or plain.startswith("Última atualização:")
        or plain.startswith("Versão:")
    )


class Builder:
    def __init__(self, doc: Document, skip_handoff: bool = True):
        self.doc = doc
        self.skip_handoff = skip_handoff
        self.block_count = 0
        self._first_body = True

    def add_title(self, html: str) -> None:
        p = self.doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=8, line=276)
        add_inline_runs(p, html, size=18, color=NAVY_RGB)
        for run in p.runs:
            run.bold = True
        self.block_count += 1

    def add_heading(self, html: str, level: int) -> None:
        p = self.doc.add_paragraph()
        if level == 2:
            set_paragraph_spacing(p, before=14, after=6, line=276)
            add_inline_runs(p, html, size=14, color=NAVY_RGB)
        else:
            set_paragraph_spacing(p, before=10, after=4, line=260)
            add_inline_runs(p, html, size=12, color=TEAL_RGB)
        for run in p.runs:
            run.bold = True
        self.block_count += 1

    def add_paragraph(self, html: str) -> None:
        if not html:
            return
        if self.skip_handoff and "handoff" in html.lower():
            return
        p = self.doc.add_paragraph()
        if looks_like_meta(html):
            set_paragraph_spacing(p, before=0, after=2, line=240)
            add_inline_runs(p, html, size=10, color=MUTED)
        elif self._first_body:
            self._first_body = False
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_paragraph_spacing(p, before=6, after=10, line=276)
            add_inline_runs(p, html, size=11, color=BODY)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_paragraph_spacing(p, before=2, after=6, line=260)
            add_inline_runs(p, html, size=10.5, color=BODY)
        self.block_count += 1

    def add_list(self, items: list[str], ordered: bool) -> None:
        for i, html in enumerate(items, 1):
            p = self.doc.add_paragraph()
            set_paragraph_spacing(p, before=1, after=2, line=260, left=5)
            prefix = f"{i}. " if ordered else "• "
            run = p.add_run(prefix)
            set_run_font(run, size=10.5, color=BODY)
            add_inline_runs(p, html, size=10.5, color=BODY)
        if items:
            spacer = self.doc.add_paragraph()
            set_paragraph_spacing(spacer, before=0, after=4, line=160)
        self.block_count += 1

    def add_hr(self) -> None:
        p = self.doc.add_paragraph()
        set_paragraph_spacing(p, before=4, after=4, line=80)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), RULE)
        pBdr.append(bottom)
        pPr.append(pBdr)
        self.block_count += 1

    def add_code(self, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        disable_autofit(table)
        set_table_fixed(table)
        set_table_width(table, USABLE_MM)
        cell = table.cell(0, 0)
        shade_cell(cell, ROW_ALT)
        set_cell_borders(cell, GRID, "4")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, before=0, after=0, line=220)
        run = p.add_run(text.strip("\n"))
        set_run_font(run, size=8, color=BODY, mono=True)
        after = self.doc.add_paragraph()
        set_paragraph_spacing(after, before=0, after=6, line=120)
        self.block_count += 1

    def add_table(self, rows: list[list[tuple[str, bool]]]) -> None:
        if not rows:
            return
        n_cols = max(len(r) for r in rows)
        padded = [r + [("", False)] * (n_cols - len(r)) for r in rows]
        first_plain = re.sub(r"<[^>]+>", "", padded[0][0][0] if padded[0] else "")
        widths = self._col_widths(n_cols, first_plain)

        table = self.doc.add_table(rows=len(padded), cols=n_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        disable_autofit(table)
        set_table_fixed(table)
        set_table_width(table, USABLE_MM)

        for r_idx, row in enumerate(padded):
            prevent_row_split(table.rows[r_idx])
            for c_idx, (html, is_head) in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.width = Mm(widths[c_idx])
                set_cell_valign(cell, "top")
                set_cell_margins(cell, top=40, bottom=40, left=50, right=50)
                set_cell_borders(cell, GRID, "4")
                if r_idx == 0:
                    shade_cell(cell, NAVY)
                elif r_idx % 2 == 0:
                    shade_cell(cell, ROW_ALT)
                p = cell.paragraphs[0]
                set_paragraph_spacing(p, before=0, after=0, line=220)
                color = WHITE if r_idx == 0 else BODY
                add_inline_runs(p, html, size=8, color=color)
                if r_idx == 0:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = WHITE

        after = self.doc.add_paragraph()
        set_paragraph_spacing(after, before=2, after=8, line=120)
        self.block_count += 1

    @staticmethod
    def _col_widths(n_cols: int, first: str) -> list[float]:
        if n_cols == 4:
            if "Método" in first:
                return [18, 58, 16, 82]
            if first.strip() == "#":
                return [10, 78, 16, 70]
            return [USABLE_MM / n_cols] * n_cols
        if n_cols == 3 and "Método" in first:
            return [18, 62, 94]
        if n_cols == 2:
            return [48, 126]
        return [USABLE_MM / n_cols] * n_cols


class MdHtmlParser(HTMLParser):
    def __init__(self, builder: Builder):
        super().__init__(convert_charrefs=True)
        self.b = builder
        self._buf: list[str] = []
        self._list_items: list[str] = []
        self._list_ordered = False
        self._table_rows: list[list[tuple[str, bool]]] = []
        self._row: list[tuple[str, bool]] = []
        self._in_header_cell = False
        self._in_pre = False

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "h3", "p", "li", "th", "td"):
            self._buf = []
            self._in_header_cell = tag == "th"
        elif tag == "pre":
            self._in_pre = True
            self._buf = []
        elif tag == "strong":
            self._buf.append("<b>")
        elif tag == "em":
            self._buf.append("<i>")
        elif tag == "code":
            if not self._in_pre:
                self._buf.append("<code>")
        elif tag == "ul":
            self._list_items = []
            self._list_ordered = False
        elif tag == "ol":
            self._list_items = []
            self._list_ordered = True
        elif tag == "table":
            self._table_rows = []
        elif tag == "tr":
            self._row = []
        elif tag == "hr":
            self.b.add_hr()
        elif tag == "br":
            self._buf.append("<br/>")

    def handle_endtag(self, tag):
        if tag == "strong":
            self._buf.append("</b>")
        elif tag == "em":
            self._buf.append("</i>")
        elif tag == "code":
            if not self._in_pre:
                self._buf.append("</code>")
        elif tag == "pre":
            raw = unescape("".join(self._buf))
            self._buf = []
            self._in_pre = False
            text = raw.strip("\n")
            if text:
                self.b.add_code(text)
        elif tag == "h1":
            self.b.add_title(self._flush())
        elif tag == "h2":
            self.b.add_heading(self._flush(), 2)
        elif tag == "h3":
            self.b.add_heading(self._flush(), 3)
        elif tag == "p":
            self.b.add_paragraph(self._flush())
        elif tag == "li":
            self._list_items.append(self._flush())
        elif tag in ("ul", "ol"):
            self.b.add_list(self._list_items, self._list_ordered)
            self._list_items = []
        elif tag in ("td", "th"):
            self._row.append((self._flush(), self._in_header_cell))
        elif tag == "tr":
            if self._row:
                self._table_rows.append(self._row)
            self._row = []
        elif tag == "table":
            self.b.add_table(self._table_rows)
            self._table_rows = []

    def handle_data(self, data):
        if not data:
            return
        if self._in_pre:
            self._buf.append(data)
        else:
            self._buf.append(
                data.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            )

    def _flush(self) -> str:
        text = "".join(self._buf).strip()
        self._buf = []
        return text


def extract_h1(md: str) -> str:
    match = re.search(r"^#\s+(.+)$", md, re.M)
    return match.group(1).strip() if match else ""


def extract_version(md: str) -> str:
    match = re.search(r"\*\*Versão:\*\*\s*([0-9.]+)", md)
    if match:
        return match.group(1)
    match = re.search(r"^Versão:\s*([0-9.]+)", md, re.M)
    return match.group(1) if match else ""


def default_header(author: str, title: str) -> str:
    shortened = title.replace(" — ", " · ", 1) if title else ""
    if not shortened:
        return author
    return f"{author} — {shortened}"


def setup_section(doc: Document, header_title: str, version: str) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(18)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(6)

    header = section.header
    header.is_linked_to_previous = False
    # Clear default empty paragraph
    for p in list(header.paragraphs):
        p._element.getparent().remove(p._element)

    htable = header.add_table(1, 1, Mm(USABLE_MM))
    disable_autofit(htable)
    set_table_fixed(htable)
    set_table_width(htable, USABLE_MM)
    cell = htable.cell(0, 0)
    shade_cell(cell, NAVY)
    set_cell_borders(cell, NAVY, "0")
    set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
    set_cell_valign(cell, "center")
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=0, after=0, line=200)
    left = p.add_run(header_title)
    set_run_font(left, size=9, color=WHITE)
    if version:
        spacer = p.add_run("\t")
        set_run_font(spacer, size=9, color=WHITE)
        right = p.add_run(f"v{version}")
        set_run_font(right, size=9, color=WHITE)
        # tab stop at right edge
        pPr = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:pos"), str(int(USABLE_MM * 56.7) - 160))
        tabs.append(tab)
        pPr.append(tabs)

    footer = section.footer
    footer.is_linked_to_previous = False
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)

    ftable = footer.add_table(1, 1, Mm(USABLE_MM))
    disable_autofit(ftable)
    set_table_fixed(ftable)
    set_table_width(ftable, USABLE_MM)
    fcell = ftable.cell(0, 0)
    shade_cell(fcell, NAVY)
    set_cell_borders(fcell, NAVY, "0")
    set_cell_margins(fcell, top=40, bottom=40, left=80, right=80)
    set_cell_valign(fcell, "center")
    fp = fcell.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(fp, before=0, after=0, line=200)
    prefix = fp.add_run("Página ")
    set_run_font(prefix, size=8, color=WHITE)
    add_page_field(fp, "PAGE")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Converte Markdown para Word no layout Poly (faixa navy, tabelas zebradas)."
    )
    parser.add_argument("input", type=Path, help="Arquivo .md de origem")
    parser.add_argument(
        "output",
        type=Path,
        nargs="?",
        help="Arquivo .docx de destino (default: mesmo nome ao lado do .md)",
    )
    parser.add_argument("--title", help="Título do documento (metadata). Default: H1 do Markdown")
    parser.add_argument("--header", help="Texto da faixa superior. Default: '{author} — {título}'")
    parser.add_argument("--subject", help="Subject do documento")
    parser.add_argument("--author", default="Poly", help="Autor / marca da faixa (default: Poly)")
    parser.add_argument("--version", help="Versão na faixa (default: campo **Versão:** do Markdown)")
    parser.add_argument(
        "--keep-handoff",
        action="store_true",
        help="Inclui parágrafos que mencionam handoff (por padrão são omitidos)",
    )
    return parser.parse_args(argv)


def build_docx(args: argparse.Namespace) -> dict:
    src = args.input.expanduser().resolve()
    if not src.is_file():
        sys.stderr.write(f"arquivo não encontrado: {src}\n")
        sys.exit(1)

    dest = args.output.expanduser().resolve() if args.output else src.with_suffix(".docx")
    dest.parent.mkdir(parents=True, exist_ok=True)

    md = src.read_text(encoding="utf-8")
    title = args.title or extract_h1(md) or src.stem
    version = args.version if args.version is not None else extract_version(md)
    header_title = args.header or default_header(args.author, title)
    subject = args.subject or title

    doc = Document()
    core = doc.core_properties
    core.title = title
    core.author = args.author
    core.subject = subject

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY

    setup_section(doc, header_title, version)
    html = markdown.markdown(md, extensions=["tables", "sane_lists", "nl2br", "fenced_code"])
    builder = Builder(doc, skip_handoff=not args.keep_handoff)
    MdHtmlParser(builder).feed(html)

    doc.save(str(dest))
    return {
        "output": str(dest),
        "title": title,
        "header": header_title,
        "version": version,
        "tables": len(doc.tables),
        "paragraphs": builder.block_count,
    }


def main() -> None:
    result = build_docx(parse_args())
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
