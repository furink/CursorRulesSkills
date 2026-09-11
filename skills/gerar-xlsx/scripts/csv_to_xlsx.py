#!/usr/bin/env python3
"""CSV (or Markdown GFM tables) → Excel with the Poly navy list layout."""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

NAVY = "1A365D"
ROW_ALT = "EDF2F7"
GRID = "A0AEC0"
MUTED = "4A5568"
WHITE = "FFFFFF"
FONT = "Calibri"

PRIO_FILLS = {
    "P0": "9B2335",
    "P1": "C47B17",
    "P2": "2E5A88",
    "P3": "5C6B73",
}

PRIO_LEGEND = {
    "P0": "Já pedido ao ERP e ainda incompleto",
    "P1": "Amplia endpoint existente",
    "P2": "Endpoint novo",
    "P3": "Relatório / opcional",
}

ILLEGAL_SHEET = re.compile(r"[\[\]:*?/\\]")


def extract_h1(md: str) -> str:
    match = re.search(r"^#\s+(.+)$", md, re.M)
    return match.group(1).strip() if match else ""


def extract_version(md: str) -> str:
    match = re.search(r"\*\*Versão:\*\*\s*([0-9.]+)", md)
    if match:
        return match.group(1)
    match = re.search(r"^Versão:\s*([0-9.]+)", md, re.M)
    return match.group(1) if match else ""


def default_header(author: str, title: str) -> str:
    shortened = title.replace(" — ", " · ", 1) if title else ""
    if not shortened:
        return author
    return f"{author} — {shortened}"


def safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = ILLEGAL_SHEET.sub(" ", (name or "Dados").strip()) or "Dados"
    cleaned = re.sub(r"\s+", " ", cleaned)[:31]
    base = cleaned
    n = 2
    while cleaned in used:
        suffix = f" ({n})"
        cleaned = (base[: 31 - len(suffix)] + suffix).strip()
        n += 1
    used.add(cleaned)
    return cleaned


def sniff_delimiter(sample: str) -> str:
    first = sample.splitlines()[0] if sample.strip() else ""
    if first.count(";") >= first.count(","):
        return ";"
    return ","


def normalize_row(row: list[str], n_cols: int) -> list[str]:
    if len(row) < n_cols:
        return row + [""] * (n_cols - len(row))
    if len(row) > n_cols:
        return row[: n_cols - 1] + [";".join(row[n_cols - 1 :])]
    return row


def read_csv_table(path: Path, delimiter: str | None) -> tuple[list[str], list[list[str]]]:
    text = path.read_text(encoding="utf-8-sig")
    delim = delimiter or sniff_delimiter(text)
    reader = csv.reader(text.splitlines(), delimiter=delim)
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        sys.stderr.write(f"CSV vazio: {path}\n")
        sys.exit(1)
    header = [c.strip() for c in rows[0]]
    data = [normalize_row([c.strip() for c in r], len(header)) for r in rows[1:]]
    return header, data


def parse_md_tables(md: str) -> list[tuple[str, list[str], list[list[str]]]]:
    lines = md.splitlines()
    tables: list[tuple[str, list[str], list[list[str]]]] = []
    last_heading = "Dados"
    i = 0
    while i < len(lines):
        line = lines[i]
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            last_heading = heading.group(2).strip()
            i += 1
            continue
        if "|" not in line:
            i += 1
            continue
        block = []
        while i < len(lines) and "|" in lines[i]:
            block.append(lines[i])
            i += 1
        parsed = _gfm_table(block)
        if parsed:
            header, data = parsed
            tables.append((last_heading, header, data))
    return tables


def _gfm_table(block: list[str]) -> tuple[list[str], list[list[str]]] | None:
    rows = []
    for raw in block:
        line = raw.strip()
        if line.startswith("|"):
            line = line[1:]
        if line.endswith("|"):
            line = line[:-1]
        cells = [c.strip() for c in line.split("|")]
        if cells and re.match(r"^:?-+:?$", cells[0].replace(" ", "")):
            continue
        if all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells if c):
            continue
        rows.append(cells)
    if not rows:
        return None
    header = rows[0]
    data = [normalize_row(r, len(header)) for r in rows[1:]]
    return header, data


def legend_from_md(md: str) -> list[tuple[str, str]]:
    tables = parse_md_tables(md)
    for name, header, data in tables:
        cols = [h.lower() for h in header]
        if "coluna" in cols and any("significado" in c or "descrição" in c or "descricao" in c for c in cols):
            idx_col = cols.index("coluna")
            idx_sig = next(
                i
                for i, c in enumerate(cols)
                if "significado" in c or "descrição" in c or "descricao" in c
            )
            pairs = []
            for row in data:
                col = row[idx_col].strip("`") if idx_col < len(row) else ""
                sig = row[idx_sig] if idx_sig < len(row) else ""
                if col:
                    pairs.append((col, sig))
            return pairs
        if "como ler" in name.lower():
            if len(header) >= 2:
                return [(r[0].strip("`"), r[1]) for r in data if r and r[0]]
    return []


def sibling_md(src: Path) -> Path | None:
    candidate = src.with_suffix(".md")
    return candidate if candidate.is_file() and candidate != src else None


def thin_border() -> Border:
    side = Side(style="thin", color=GRID)
    return Border(left=side, right=side, top=side, bottom=side)


def navy_fill() -> PatternFill:
    return PatternFill("solid", fgColor=NAVY)


def zebra_fill() -> PatternFill:
    return PatternFill("solid", fgColor=ROW_ALT)


def write_data_sheet(
    ws: Worksheet,
    header: list[str],
    data: list[list[str]],
    print_header: str,
    version: str,
) -> None:
    border = thin_border()
    header_font = Font(name=FONT, bold=True, color=WHITE, size=11)
    cell_font = Font(name=FONT, size=10)
    prio_font = Font(name=FONT, bold=True, color=WHITE, size=11)
    wrap = Alignment(vertical="center", wrap_text=True)
    wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    prio_idx = None
    for i, name in enumerate(header):
        if name.lower() in ("prioridade", "priority", "prio"):
            prio_idx = i
            break

    for col, name in enumerate(header, start=1):
        cell = ws.cell(1, col, name)
        cell.font = header_font
        cell.fill = navy_fill()
        cell.alignment = wrap_center
        cell.border = border

    for r_idx, row in enumerate(data, start=2):
        zebra = r_idx % 2 == 0
        prio_val = row[prio_idx].strip().upper() if prio_idx is not None and prio_idx < len(row) else ""
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.alignment = wrap
            cell.border = border
            cell.font = cell_font
            if prio_idx is not None and c_idx == prio_idx + 1 and prio_val in PRIO_FILLS:
                cell.fill = PatternFill("solid", fgColor=PRIO_FILLS[prio_val])
                cell.font = prio_font
                cell.alignment = wrap_center
            elif zebra:
                cell.fill = zebra_fill()
        ws.row_dimensions[r_idx].height = 32

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(header))
    last_row = 1 + len(data)
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    for col_idx, name in enumerate(header, start=1):
        samples = [name] + [r[col_idx - 1] for r in data[:80] if col_idx - 1 < len(r)]
        width = max((len(s) for s in samples), default=10)
        width = min(max(width + 2, 10), 48)
        if col_idx == 1:
            width = min(max(width, 12), 14)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:1"
    ws.page_setup.horizontalCentered = True
    ws.oddHeader.left.text = print_header
    if version:
        ws.oddHeader.right.text = f"v{version}"
    ws.oddFooter.right.text = "Pág. &P / &N"
    ws.sheet_properties.tabColor = NAVY


def write_legend_sheet(
    ws: Worksheet,
    title: str,
    version: str,
    columns: list[str],
    column_meanings: list[tuple[str, str]],
    has_priority: bool,
) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    meta = f"Versão {version}" if version else ""
    today = date.today().strftime("%d-%m-%Y")
    ws["A2"] = " — ".join(p for p in (meta, today) if p)
    ws["A2"].font = Font(name=FONT, size=11, color=MUTED)
    ws.row_dimensions[1].height = 22

    meanings = {k.lower().strip("`"): v for k, v in column_meanings}
    rows = [("Coluna", "Significado")]
    for col in columns:
        rows.append((col, meanings.get(col.lower(), meanings.get(col.lower().strip("`"), ""))))

    if has_priority:
        rows.append(("", ""))
        rows.append(("Prioridade", "Uso"))
        for code, text in PRIO_LEGEND.items():
            rows.append((code, text))

    border = thin_border()
    header_font = Font(name=FONT, bold=True, color=WHITE)
    start = 4
    for offset, (a, b) in enumerate(rows):
        r = start + offset
        c1 = ws.cell(r, 1, a)
        c2 = ws.cell(r, 2, b)
        c1.alignment = Alignment(vertical="center", wrap_text=True)
        c2.alignment = Alignment(vertical="center", wrap_text=True)
        is_header = (a, b) in (("Coluna", "Significado"), ("Prioridade", "Uso"))
        if is_header:
            for cell in (c1, c2):
                cell.font = header_font
                cell.fill = navy_fill()
                cell.border = border
        else:
            c1.font = Font(name=FONT, bold=True)
            c2.font = Font(name=FONT, size=11)
            c1.border = border
            c2.border = border
            if a in PRIO_FILLS:
                c1.fill = PatternFill("solid", fgColor=PRIO_FILLS[a])
                c1.font = Font(name=FONT, bold=True, color=WHITE)
                c1.alignment = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 95
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "C47B17"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Converte CSV ou tabelas Markdown para Excel no layout Poly (navy, zebrado, Legenda)."
    )
    parser.add_argument("input", type=Path, help="Arquivo .csv ou .md de origem")
    parser.add_argument(
        "output",
        type=Path,
        nargs="?",
        help="Arquivo .xlsx de destino (default: mesmo nome ao lado da origem)",
    )
    parser.add_argument("--title", help="Título (metadata e aba Legenda). Default: H1 / stem")
    parser.add_argument("--header", help="Texto do cabeçalho de impressão. Default: '{author} — {título}'")
    parser.add_argument("--subject", help="Subject do arquivo")
    parser.add_argument("--author", default="Poly", help="Autor / marca (default: Poly)")
    parser.add_argument("--version", help="Versão (default: campo **Versão:** do Markdown)")
    parser.add_argument("--sheet", help="Nome da aba de dados (CSV). Default: Dados")
    parser.add_argument("--delimiter", help="Delimitador do CSV (default: detecta ; ou ,)")
    parser.add_argument("--no-legend", action="store_true", help="Não gera a aba Legenda")
    return parser.parse_args(argv)


def build_xlsx(args: argparse.Namespace) -> dict:
    src = args.input.expanduser().resolve()
    if not src.is_file():
        sys.stderr.write(f"arquivo não encontrado: {src}\n")
        sys.exit(1)

    dest = args.output.expanduser().resolve() if args.output else src.with_suffix(".xlsx")
    dest.parent.mkdir(parents=True, exist_ok=True)

    companion = sibling_md(src)
    md_text = ""
    if src.suffix.lower() == ".md":
        md_text = src.read_text(encoding="utf-8")
    elif companion:
        md_text = companion.read_text(encoding="utf-8")

    title = args.title or (extract_h1(md_text) if md_text else "") or src.stem
    version = args.version if args.version is not None else (extract_version(md_text) if md_text else "")
    header_title = args.header or default_header(args.author, title)
    subject = args.subject or title
    meanings = legend_from_md(md_text) if md_text else []

    wb = Workbook()
    used_names: set[str] = set()
    sheets_meta: list[dict] = []
    first = True

    if src.suffix.lower() == ".md":
        tables = parse_md_tables(md_text)
        if not tables:
            sys.stderr.write("nenhuma tabela Markdown encontrada\n")
            sys.exit(1)
        for heading, header, data in tables:
            name = safe_sheet_name(heading, used_names)
            ws = wb.active if first else wb.create_sheet(name)
            if first:
                ws.title = name
                first = False
            write_data_sheet(ws, header, data, header_title, version)
            sheets_meta.append({"name": name, "rows": len(data), "header": header})
    else:
        header, data = read_csv_table(src, args.delimiter)
        name = safe_sheet_name(args.sheet or "Dados", used_names)
        ws = wb.active
        ws.title = name
        write_data_sheet(ws, header, data, header_title, version)
        sheets_meta.append({"name": name, "rows": len(data), "header": header})

    if not args.no_legend and sheets_meta:
        legend_name = safe_sheet_name("Legenda", used_names)
        legend = wb.create_sheet(legend_name, 1)
        primary = sheets_meta[0]
        has_prio = any(h.lower() in ("prioridade", "priority", "prio") for h in primary["header"])
        write_legend_sheet(legend, title, version, primary["header"], meanings, has_prio)
        sheets_meta.insert(1, {"name": legend_name, "rows": 0, "header": []})

    props = wb.properties
    props.title = title
    props.creator = args.author
    props.subject = subject

    wb.save(str(dest))
    total_rows = sum(s["rows"] for s in sheets_meta)
    return {
        "output": str(dest),
        "sheets": [s["name"] for s in sheets_meta],
        "rows": total_rows,
        "title": title,
        "header": header_title,
        "version": version,
    }


def main() -> None:
    result = build_xlsx(parse_args())
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
